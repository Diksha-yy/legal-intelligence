"""Official legal PDF collector. Does not bypass India Code viewer or access restrictions."""
import argparse,csv,hashlib,json,re,sys
from datetime import datetime,timezone
from pathlib import Path
from urllib.parse import urlparse
import requests
from pypdf import PdfReader
from openpyxl import load_workbook

BASE=Path(__file__).resolve().parent
ALLOWED={'indiacode.gov.in','www.indiacode.gov.in','indiacode.nic.in','www.indiacode.nic.in','cdnbbsr.s3waas.gov.in'}
HEADERS=['document_id','law_id','domain','title','document_type','jurisdiction','year','source_page_url','pdf_url','publication_date','effective_date','version_date','legal_status','verification_status','download_status']

def read_manifest(path):
 if path.suffix.lower()=='.xlsx':
  wb=load_workbook(path);ws=wb['Sources'] if 'Sources' in wb else wb.active
  rows=list(ws.values); cols=[str(c or '').strip() for c in rows[0]]
  return [dict(zip(cols,[str(v).strip() if v is not None else '' for v in r])) for r in rows[1:] if any(v is not None for v in r)], ('xlsx',wb,ws,cols)
 with path.open(encoding='utf-8-sig',newline='') as f:
  dr=csv.DictReader(f);return list(dr),('csv',None,None,dr.fieldnames)

def official(url):
 p=urlparse(url)
 return p.scheme=='https' and p.hostname in ALLOWED and not p.username and not p.password

def inspect_pdf(path):
 with path.open('rb') as f:
  if f.read(5)!=b'%PDF-':raise ValueError('Not a PDF signature')
 digest=hashlib.sha256(path.read_bytes()).hexdigest()
 reader=PdfReader(str(path),strict=False)
 if len(reader.pages)<1:raise ValueError('PDF has no pages')
 sample=' '.join((reader.pages[i].extract_text() or '') for i in range(min(3,len(reader.pages))))
 return {'sha256':digest,'pages':len(reader.pages),'sample_text':sample[:2500],'ocr_required':len(sample.strip())<120}

def download_pdf(url,path):
 if not official(url):raise ValueError('Source not in official allowlist')
 if '/pdf-viewer/' in urlparse(url).path:raise ValueError('Viewer page, not a direct PDF')
 with requests.get(url,timeout=(12,45),stream=True,headers={'User-Agent':'LegalIntelligenceAcademicCollector/1.0'},allow_redirects=True) as r:
  r.raise_for_status()
  if not official(r.url):raise ValueError('Redirected outside official allowlist')
  length=0
  tmp=path.with_suffix('.part')
  try:
   with tmp.open('wb') as f:
    for block in r.iter_content(65536):
     if block:
      length+=len(block)
      if length>35*1024*1024:raise ValueError('PDF exceeds 35MB safety cap')
      f.write(block)
   with tmp.open('rb') as f:
    if f.read(5)!=b'%PDF-':raise ValueError('Response is not PDF (perhaps HTML)')
   tmp.replace(path)
  finally:tmp.unlink(missing_ok=True)

def main():
 p=argparse.ArgumentParser();p.add_argument('--input',default='sources.xlsx');p.add_argument('--extract',action='store_true');p.add_argument('--no-download',action='store_true');a=p.parse_args()
 manifest=Path(a.input);manifest=manifest if manifest.is_absolute() else BASE/manifest
 rows,meta=read_manifest(manifest)
 if not set(HEADERS).issubset(meta[3]):sys.exit('Manifest is missing required columns')
 out=BASE/'downloads';out.mkdir(exist_ok=True);report=[];seen={};chunks=[]
 for row in rows:
  doc=row['document_id'];domain=row['domain'].lower();src=row['source_page_url'];url=row['pdf_url']
  result={'document_id':doc,'title':row['title'],'source_page_url':src,'pdf_url':url,'checked_at':datetime.now(timezone.utc).isoformat(),'status':'','verification':'pending','sha256':'','pages':'','ocr_required':'','note':''}
  try:
   if not re.fullmatch(r'[A-Z]+_\d{3}',doc):raise ValueError('Invalid document_id')
   if domain not in {'rental','divorce','property','employment'}:raise ValueError('Invalid domain')
   if src and not official(src):raise ValueError('Source page is not official allowlist')
   directory=out/domain;directory.mkdir(parents=True,exist_ok=True);dest=directory/(doc+'.pdf')
   if not dest.exists():
    if not url or '/pdf-viewer/' in urlparse(url).path:
     result['status']='manual_required';result['note']='Open source_page_url, use viewer Download button, save as downloads/'+domain+'/'+doc+'.pdf'
    elif a.no_download:result['status']='pending';result['note']='Download disabled'
    else:download_pdf(url,dest)
   if dest.exists():
    info=inspect_pdf(dest);result.update({'status':'downloaded','sha256':info['sha256'],'pages':info['pages'],'ocr_required':info['ocr_required']})
    result['verification']='needs_review'
    if info['sha256'] in seen:result['note']='Duplicate bytes of '+seen[info['sha256']]
    else:seen[info['sha256']]=doc
    if info['ocr_required']:result['note']+=' Text extraction low; OCR required.'
    if a.extract:
     reader=PdfReader(str(dest),strict=False)
     for page_no,page in enumerate(reader.pages,1):
      try:txt=page.extract_text() or ''
      except Exception:txt=''
      if txt.strip():chunks.append({'document_id':doc,'law_id':row['law_id'],'domain':domain,'jurisdiction':row['jurisdiction'],'page':page_no,'text':txt,'sha256':info['sha256'],'legal_status':row['legal_status'],'verification_status':'needs_review','source_page_url':src})
  except Exception as e:result['status']='failed';result['note']=str(e)[:350]
  row['download_status']=result['status'];row['verification_status']=result['verification']
  report.append(result);print(f"{doc}: {result['status']} {result['note']}")
 with (BASE/'collection_report.csv').open('w',newline='',encoding='utf-8-sig') as f:
  writer=csv.DictWriter(f,fieldnames=list(report[0]) if report else []);writer.writeheader();writer.writerows(report)
 if a.extract:
  (BASE/'processed').mkdir(exist_ok=True)
  with (BASE/'processed'/'pages.jsonl').open('w',encoding='utf-8') as f:
   for item in chunks:f.write(json.dumps(item,ensure_ascii=False)+'\n')
 # Avoid rewriting the user's original input: create separate output workbook/CSV.
 output=BASE/('sources_updated.xlsx' if meta[0]=='xlsx' else 'sources_updated.csv')
 if meta[0]=='xlsx':
  wb,ws,cols=meta[1:]
  for index,row in enumerate(rows,2):
   for col in ['download_status','verification_status']:
    ws.cell(index,cols.index(col)+1,row[col])
  wb.save(output)
 else:
  with output.open('w',encoding='utf-8-sig',newline='') as f:
   writer=csv.DictWriter(f,fieldnames=meta[3]);writer.writeheader();writer.writerows(rows)
 print('Report:',BASE/'collection_report.csv');print('Updated manifest:',output)

if __name__=='__main__':main()
